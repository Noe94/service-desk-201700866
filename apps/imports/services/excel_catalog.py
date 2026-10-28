from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries


HEADERS = [
    "COD.N1",
    "SERVICIO - Nivel 1",
    "COD.N2",
    "SERVICIO - Nivel 2",
    "ACTIVO",
    "CLASE DE SERVICIO",
    "CRITICIDAD",
    "TIPO DE SERVICIO",
    "Descripción",
    "Métrica",
    "Minimo",
    "Maximo",
]


def text(value):
    return "" if value is None else str(value).strip()


class CatalogWorkbookReader:
    def __init__(self, path, sheet_name="Servicios Externos"):
        self.path = Path(path)

        if not self.path.exists():
            raise ValueError(f"No existe el archivo: {self.path}")

        self.workbook = load_workbook(
            self.path,
            data_only=True,
            read_only=False,
        )

        if sheet_name not in self.workbook.sheetnames:
            raise ValueError(
                f"No existe la hoja '{sheet_name}'. "
                f"Disponibles: {', '.join(self.workbook.sheetnames)}"
            )

        self.sheet_name = sheet_name
        self.ws = self.workbook[sheet_name]
        self.merged = self._build_merged_lookup()

    def _build_merged_lookup(self):
        lookup = {}

        for merged_range in self.ws.merged_cells.ranges:
            min_col, min_row, max_col, max_row = range_boundaries(
                str(merged_range)
            )

            info = (min_row, min_col, max_row, max_col)

            for row in range(min_row, max_row + 1):
                for col in range(min_col, max_col + 1):
                    lookup[(row, col)] = info

        return lookup

    def effective_value(self, row, col):
        physical_value = self.ws.cell(row, col).value

        if physical_value is not None:
            return physical_value

        info = self.merged.get((row, col))

        if info is None:
            return None

        min_row, min_col, _, _ = info

        return self.ws.cell(min_row, min_col).value

    def end_row(self, row, *columns):
        result = row

        for col in columns:
            info = self.merged.get((row, col))

            if info:
                result = max(result, info[2])

        return result

    def source_range(self, row, start_col, end_col, *anchors):
        final_row = self.end_row(row, *anchors)

        return (
            f"{get_column_letter(start_col)}{row}:"
            f"{get_column_letter(end_col)}{final_row}",
            final_row,
        )

    def find_header_row(self):
        search_limit = min(self.ws.max_row, 100)

        for row in range(1, search_limit + 1):
            values = [
                text(self.ws.cell(row, col).value)
                for col in range(1, 13)
            ]

            if values == HEADERS:
                return row

        raise ValueError(
            "No se encontró la fila de encabezados A:L esperada."
        )

    def find_options_row(self, start_row):
        for row in range(start_row, self.ws.max_row + 1):
            values = [
                text(self.ws.cell(row, col).value).upper()
                for col in range(5, 9)
            ]

            if values == ["OPCIONES"] * 4:
                return row

        raise ValueError(
            "No se encontró la sección OPCIONES del machote."
        )

    def read_options(self, options_row):
        columns = {
            "activo": 5,
            "clase": 6,
            "criticidad": 7,
            "tipo": 8,
        }

        result = {}

        for name, col in columns.items():
            values = []

            for row in range(
                options_row + 1,
                self.ws.max_row + 1,
            ):
                value = text(
                    self.ws.cell(row, col).value
                )

                if value and value not in values:
                    values.append(value)

            result[name] = values

        return result

    def parse(self):
        header_row = self.find_header_row()
        options_row = self.find_options_row(header_row + 1)
        options = self.read_options(options_row)

        logical_rows = [
            row
            for row in range(header_row + 1, options_row)
            if text(self.ws.cell(row, 3).value)
        ]

        observations = []
        n1 = {}

        # Primera pasada:
        # obtiene N1 y detecta conflictos de nombres.
        for row in logical_rows:
            code = text(self.effective_value(row, 1))
            name = text(self.effective_value(row, 2))

            if not code:
                continue

            source, final_row = self.source_range(
                row,
                1,
                2,
                1,
                2,
            )

            if code not in n1:
                n1[code] = {
                    "codigo": code,
                    "nombre": name,
                    "fila_inicio": row,
                    "fila_fin": final_row,
                    "rango_origen": source,
                }

            elif name and name != n1[code]["nombre"]:
                observations.append(
                    {
                        "tipo": "CONFLICTING_N1_NAME",
                        "nivel": "N1",
                        "codigo": code,
                        "fila_inicio": row,
                        "fila_fin": row,
                        "rango_origen": f"A{row}:B{row}",
                        "mensaje": (
                            "El código N1 presenta nombres distintos. "
                            "Se conserva la primera aparición válida."
                        ),
                        "valores_originales": {
                            "canonico": n1[code]["nombre"],
                            "conflictivo": name,
                        },
                        "valor_canonico": n1[code]["nombre"],
                    }
                )

        n2 = []

        # Segunda pasada:
        # cada código físico de la columna C inicia un N2 lógico.
        for row in logical_rows:
            code = text(
                self.ws.cell(row, 3).value
            )

            name = text(
                self.effective_value(row, 4)
            )

            parent_code = text(
                self.effective_value(row, 1)
            )

            parent_name = text(
                self.effective_value(row, 2)
            )

            transformations = {}
            requires_review = False

            # Si N1 está vacío y NO proviene de una combinación,
            # se aplica una regla explícita basada en el código N2.
            if not parent_code:
                derived = (
                    code.rsplit(".", 1)[0]
                    if "." in code
                    else ""
                )

                if derived in n1:
                    parent_code = derived
                    parent_name = n1[derived]["nombre"]

                    transformations["nivel1"] = {
                        "regla": "DERIVED_FROM_N2_CODE",
                        "codigo_derivado": derived,
                    }

                    observations.append(
                        {
                            "tipo": "PARENT_DERIVED_FROM_N2_CODE",
                            "nivel": "N2",
                            "codigo": code,
                            "fila_inicio": row,
                            "fila_fin": row,
                            "rango_origen": f"A{row}:D{row}",
                            "mensaje": (
                                "La fila no tiene N1 y no pertenece "
                                "a una combinación. El padre se derivó "
                                "del prefijo del código N2."
                            ),
                            "valores_originales": {
                                "codigo_n1": "",
                                "codigo_n2": code,
                            },
                            "valor_canonico": derived,
                        }
                    )

                else:
                    requires_review = True

                    observations.append(
                        {
                            "tipo": "MISSING_PARENT",
                            "nivel": "N2",
                            "codigo": code,
                            "fila_inicio": row,
                            "fila_fin": row,
                            "rango_origen": f"A{row}:D{row}",
                            "mensaje": (
                                "No fue posible determinar el N1."
                            ),
                            "valores_originales": {
                                "codigo_n2": code,
                            },
                            "valor_canonico": None,
                        }
                    )

            raw = {
                "activo": text(
                    self.ws.cell(row, 5).value
                ),
                "clase": text(
                    self.ws.cell(row, 6).value
                ),
                "criticidad": text(
                    self.ws.cell(row, 7).value
                ),
                "tipo": text(
                    self.ws.cell(row, 8).value
                ),
                "metrica": self.effective_value(
                    row,
                    10,
                ),
            }

            missing = [
                key
                for key, value in raw.items()
                if not text(value)
            ]

            if missing:
                requires_review = True

                observations.append(
                    {
                        "tipo": "INCOMPLETE_ATTRIBUTES",
                        "nivel": "N2",
                        "codigo": code,
                        "fila_inicio": row,
                        "fila_fin": row,
                        "rango_origen": f"E{row}:L{row}",
                        "mensaje": (
                            "Hay atributos sin información; "
                            "se conservarán como desconocidos."
                        ),
                        "valores_originales": {
                            "campos_sin_informacion": missing,
                        },
                        "valor_canonico": None,
                    }
                )

            for key in (
                "activo",
                "clase",
                "criticidad",
                "tipo",
            ):
                value = raw[key]

                if value and value not in options[key]:
                    requires_review = True

                    observations.append(
                        {
                            "tipo": "VALUE_OUTSIDE_OPTIONS",
                            "nivel": "N2",
                            "codigo": code,
                            "fila_inicio": row,
                            "fila_fin": row,
                            "rango_origen": f"{key}@{row}",
                            "mensaje": (
                                f"El valor de {key} no aparece "
                                "en la lista de opciones."
                            ),
                            "valores_originales": {
                                "campo": key,
                                "valor": value,
                            },
                            "valor_canonico": None,
                        }
                    )

            source, final_row = self.source_range(
                row,
                3,
                12,
                3,
                4,
            )

            n2.append(
                {
                    "codigo": code,
                    "nombre": name,
                    "nivel1_codigo": parent_code,
                    "nivel1_nombre": parent_name,
                    "activo_original": raw["activo"],
                    "clase": raw["clase"],
                    "criticidad": raw["criticidad"],
                    "tipo": raw["tipo"],
                    "descripcion": self.ws.cell(
                        row,
                        9,
                    ).value,
                    "metrica": raw["metrica"],
                    "minimo": self.ws.cell(
                        row,
                        11,
                    ).value,
                    "maximo": self.ws.cell(
                        row,
                        12,
                    ).value,
                    "fila_inicio": row,
                    "fila_fin": final_row,
                    "rango_origen": source,
                    "transformaciones": transformations,
                    "requiere_revision": requires_review,
                }
            )

        return {
            "hoja": self.sheet_name,
            "fila_encabezado": header_row,
            "fila_opciones": options_row,
            "opciones": options,
            "nivel1": list(n1.values()),
            "nivel2": n2,
            "observaciones": observations,
        }
